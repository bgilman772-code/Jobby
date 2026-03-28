import os
import uuid
import io
import time
import zipfile
import threading
import webbrowser
from dotenv import load_dotenv
load_dotenv()

from flask import (Flask, render_template, request, redirect, url_for,
                   send_file, flash, jsonify, session)
from werkzeug.utils import secure_filename
from werkzeug.security import generate_password_hash, check_password_hash
import requests
import json
import smtplib
from email.message import EmailMessage
from cryptography.fernet import Fernet, InvalidToken
import sqlite3
from datetime import datetime, timedelta, timezone

from resume_parser import extract_text, analyze_resume
from job_matcher import apply_filters, score_and_categorize_jobs, CATEGORIES, estimate_salaries_from_comparables
from job_ranker import rank_jobs
from cover_letter_generator import tailor_cover_letter, generate_resume_feedback
from auto_filler import PENDING_APPLICATIONS, start_fill
from job_sources import fetch_all_jobs, cache_age_hours, clear_source_cache

# ── App setup ──────────────────────────────────────────────────────────────────
UPLOAD_FOLDER = os.path.join(os.path.dirname(__file__), 'uploads')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

# Stable secret key — generated once and stored so sessions survive restarts.
_SECRET_KEY_FILE = os.path.join(os.path.dirname(__file__), '.flask_secret')
if os.environ.get('SECRET_KEY'):
    app.secret_key = os.environ['SECRET_KEY'].encode()
elif os.path.exists(_SECRET_KEY_FILE):
    with open(_SECRET_KEY_FILE, 'rb') as _f:
        app.secret_key = _f.read()
else:
    _k = os.urandom(32)
    with open(_SECRET_KEY_FILE, 'wb') as _f:
        _f.write(_k)
    app.secret_key = _k

# ── Disk-backed job cache so results survive server restarts ──────────────────
_JOB_CACHE_DIR = os.path.join(os.path.dirname(__file__), '.jobcache')
os.makedirs(_JOB_CACHE_DIR, exist_ok=True)

class _DiskCache(dict):
    """dict subclass that also writes/reads each key to a JSON file."""
    def __setitem__(self, key, value):
        super().__setitem__(key, value)
        try:
            path = os.path.join(_JOB_CACHE_DIR, f'{key}.json')
            with open(path, 'w', encoding='utf-8') as f:
                json.dump(value, f)
        except Exception:
            pass

    def get(self, key, default=None):
        if key in self:
            return self[key]
        # Try loading from disk
        try:
            path = os.path.join(_JOB_CACHE_DIR, f'{key}.json')
            if os.path.exists(path):
                with open(path, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                super().__setitem__(key, data)   # populate memory cache
                return data
        except Exception:
            pass
        return default

JOB_CACHE = _DiskCache()   # cache_key -> list[job]
MATCH_TASKS: dict = {}     # task_id   -> {status, jobs, profile, cache_key, error}
BATCH_SESSIONS: dict = {}  # batch_id  -> {jobs, extra_info, statuses, current, done}

DB_PATH = os.path.join(os.path.dirname(__file__), 'jobs.db')

# Legacy flat-file paths (used as fallback / before login system existed)
PROFILE_PATH        = os.path.join(UPLOAD_FOLDER, 'profile.json')
PARSED_PROFILE_PATH = os.path.join(UPLOAD_FOLDER, 'profile_parsed.json')
RESUME_TEXT_PATH    = os.path.join(UPLOAD_FOLDER, 'resume_text.txt')


# ── Per-user path helpers ────────────────────────────────────────────────────────

def get_user_upload_folder(user_id=None) -> str:
    """Return the upload folder for a given user_id (or the shared folder if None)."""
    if user_id:
        folder = os.path.join(UPLOAD_FOLDER, f'user_{user_id}')
        os.makedirs(folder, exist_ok=True)
        return folder
    return UPLOAD_FOLDER


def _current_user_id():
    """Return the logged-in user's id from the Flask session, or None."""
    return session.get('user_id')


def get_profile_path(user_id=None) -> str:
    return os.path.join(get_user_upload_folder(user_id), 'profile.json')


def get_parsed_profile_path(user_id=None) -> str:
    return os.path.join(get_user_upload_folder(user_id), 'profile_parsed.json')


def get_resume_text_path(user_id=None) -> str:
    return os.path.join(get_user_upload_folder(user_id), 'resume_text.txt')

DEFAULT_TEMPLATE = (
    "Dear {company},\n\n"
    "I am writing to apply for the {title} role. {intro}\n\n"
    "I believe my background as {name} makes me a strong fit. {skills}\n\n"
    "Sincerely,\n{name}\n"
)


# ── Database helpers ────────────────────────────────────────────────────────────

def init_db():
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    # Users table
    c.execute('''
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            email TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            created_at TEXT NOT NULL
        )
    ''')
    # Scraped jobs table
    c.execute('''
        CREATE TABLE IF NOT EXISTS jobs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            source TEXT, url TEXT, title TEXT, company TEXT,
            description TEXT, location TEXT, remote INTEGER,
            salary TEXT, date_posted TEXT, discovered_at TEXT,
            raw TEXT, dedup_hash TEXT
        )
    ''')
    c.execute('CREATE UNIQUE INDEX IF NOT EXISTS jobs_url_idx ON jobs(url)')
    c.execute('CREATE UNIQUE INDEX IF NOT EXISTS jobs_hash_idx ON jobs(dedup_hash)')
    # Application tracker table
    c.execute('''
        CREATE TABLE IF NOT EXISTS applications (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER DEFAULT 0,
            job_url TEXT,
            job_title TEXT,
            company TEXT,
            salary TEXT,
            location TEXT,
            remote INTEGER DEFAULT 0,
            ranking_score INTEGER DEFAULT 0,
            match_score INTEGER DEFAULT 0,
            category TEXT DEFAULT 'Other',
            status TEXT DEFAULT 'saved',
            created_at TEXT,
            applied_at TEXT,
            cover_letter TEXT,
            notes TEXT,
            qualification_tier TEXT DEFAULT ''
        )
    ''')
    # Migrations for existing DBs
    for col_def in [
        'ALTER TABLE applications ADD COLUMN qualification_tier TEXT DEFAULT ""',
        'ALTER TABLE applications ADD COLUMN user_id INTEGER DEFAULT 0',
    ]:
        try:
            c.execute(col_def)
        except Exception:
            pass
    conn.commit()
    conn.close()


init_db()


# ── Auth helpers ─────────────────────────────────────────────────────────────────

_PUBLIC_ENDPOINTS = {'login', 'register', 'static'}


@app.before_request
def require_login():
    if request.endpoint and request.endpoint not in _PUBLIC_ENDPOINTS:
        if 'user_id' not in session:
            return redirect(url_for('login'))


def get_user_by_id(user_id: int):
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('SELECT id, username, email FROM users WHERE id=?', (user_id,))
    row = c.fetchone()
    conn.close()
    if row:
        return {'id': row[0], 'username': row[1], 'email': row[2]}
    return None


def get_user_by_username(username: str):
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('SELECT id, username, email, password_hash FROM users WHERE username=?', (username,))
    row = c.fetchone()
    conn.close()
    if row:
        return {'id': row[0], 'username': row[1], 'email': row[2], 'password_hash': row[3]}
    return None


def create_user(username: str, email: str, password: str) -> int:
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute(
        'INSERT INTO users (username, email, password_hash, created_at) VALUES (?,?,?,?)',
        (username, email, generate_password_hash(password),
         datetime.now(timezone.utc).isoformat())
    )
    user_id = c.lastrowid
    conn.commit()
    conn.close()
    return user_id


# ── Auth routes ──────────────────────────────────────────────────────────────────

@app.route('/login', methods=['GET', 'POST'])
def login():
    if 'user_id' in session:
        return redirect(url_for('index'))
    error = None
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        user = get_user_by_username(username)
        if user and check_password_hash(user['password_hash'], password):
            session['user_id'] = user['id']
            session['username'] = user['username']
            return redirect(url_for('index'))
        error = 'Invalid username or password.'
    return render_template('login.html', error=error)


@app.route('/register', methods=['GET', 'POST'])
def register():
    if 'user_id' in session:
        return redirect(url_for('index'))
    error = None
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        email    = request.form.get('email', '').strip()
        password = request.form.get('password', '')
        confirm  = request.form.get('confirm', '')
        if not username or not email or not password:
            error = 'All fields are required.'
        elif password != confirm:
            error = 'Passwords do not match.'
        elif len(password) < 6:
            error = 'Password must be at least 6 characters.'
        elif get_user_by_username(username):
            error = 'Username already taken.'
        else:
            try:
                user_id = create_user(username, email, password)
                session['user_id'] = user_id
                session['username'] = username
                flash(f'Welcome, {username}! Your account is ready.')
                return redirect(url_for('index'))
            except Exception as e:
                error = f'Registration failed: {e}'
    return render_template('register.html', error=error)


@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))


def query_jobs(filters: dict):
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    sql = ('SELECT id,source,url,title,company,description,location,remote,'
           'salary,date_posted,discovered_at FROM jobs WHERE 1=1')
    params = []
    if filters.get('company'):
        sql += ' AND company LIKE ?'
        params.append(f"%{filters['company']}%")
    if filters.get('title'):
        sql += ' AND title LIKE ?'
        params.append(f"%{filters['title']}%")
    if filters.get('location'):
        sql += ' AND location LIKE ?'
        params.append(f"%{filters['location']}%")
    if filters.get('remote') is not None:
        sql += ' AND remote = ?'
        params.append(1 if filters['remote'] else 0)
    if filters.get('since_days'):
        cutoff = (datetime.now(timezone.utc) - timedelta(days=int(filters['since_days']))).isoformat()
        sql += ' AND (date_posted >= ? OR discovered_at >= ?)'
        params.extend([cutoff, cutoff])
    sql += ' ORDER BY discovered_at DESC'
    c.execute(sql, params)
    rows = c.fetchall()
    conn.close()
    cols = ['id', 'source', 'url', 'title', 'company', 'description',
            'location', 'remote', 'salary', 'date_posted', 'discovered_at']
    return [dict(zip(cols, r)) for r in rows]


def save_application_record(job: dict, status: str = 'saved', cover_letter: str = '',
                            user_id=None):
    uid = user_id or _current_user_id() or 0
    job_url = job.get('url') or job.get('job_url', '')
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('''
        INSERT OR IGNORE INTO applications
          (user_id, job_url, job_title, company, salary, location, remote,
           ranking_score, match_score, category, status, created_at,
           cover_letter, qualification_tier)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)
    ''', (
        uid,
        job_url,
        job.get('title') or job.get('job_title', ''),
        job.get('company_name') or job.get('company', ''),
        job.get('salary', ''),
        job.get('candidate_required_location') or job.get('location', ''),
        1 if job.get('is_remote') or job.get('remote') else 0,
        job.get('ranking_score', 0),
        job.get('match_score', 0),
        job.get('category', 'Other'),
        status,
        datetime.now(timezone.utc).isoformat(),
        cover_letter,
        job.get('qualification_tier', ''),
    ))
    if status == 'applied':
        c.execute('UPDATE applications SET status=?, applied_at=? WHERE job_url=? AND user_id=?',
                  (status, datetime.now(timezone.utc).isoformat(), job_url, uid))
    conn.commit()
    conn.close()


def get_applications(status_filter: str = None, user_id=None):
    uid = user_id or _current_user_id() or 0
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    sql = 'SELECT * FROM applications WHERE user_id = ?'
    params = [uid]
    if status_filter:
        sql += ' AND status = ?'
        params.append(status_filter)
    sql += ' ORDER BY created_at DESC'
    c.execute(sql, params)
    rows = c.fetchall()
    cols = [d[0] for d in c.description]
    conn.close()
    return [dict(zip(cols, r)) for r in rows]


def update_application_status(app_id: int, status: str, notes: str = '', user_id=None):
    uid = user_id or _current_user_id() or 0
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    if status == 'applied':
        c.execute('UPDATE applications SET status=?, notes=?, applied_at=? WHERE id=? AND user_id=?',
                  (status, notes, datetime.now(timezone.utc).isoformat(), app_id, uid))
    else:
        c.execute('UPDATE applications SET status=?, notes=? WHERE id=? AND user_id=?',
                  (status, notes, app_id, uid))
    conn.commit()
    conn.close()


# ── Profile helpers ─────────────────────────────────────────────────────────────

def save_profile(data: dict, user_id=None):
    path = get_profile_path(user_id or _current_user_id())
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(data, f)


def load_profile(user_id=None) -> dict:
    path = get_profile_path(user_id or _current_user_id())
    if os.path.exists(path):
        with open(path, 'r', encoding='utf-8') as f:
            return json.load(f)
    return {}


def load_parsed_profile(user_id=None) -> dict:
    path = get_parsed_profile_path(user_id or _current_user_id())
    if os.path.exists(path):
        with open(path, 'r', encoding='utf-8') as f:
            return json.load(f)
    return {}


def save_parsed_profile(data: dict, user_id=None):
    path = get_parsed_profile_path(user_id or _current_user_id())
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(data, f)


def load_resume_text(user_id=None) -> str:
    path = get_resume_text_path(user_id or _current_user_id())
    if os.path.exists(path):
        with open(path, 'r', encoding='utf-8') as f:
            return f.read()
    return ''


# ── SMTP helpers ────────────────────────────────────────────────────────────────

SMTP_CONFIG_PATH = os.path.join(UPLOAD_FOLDER, 'smtp_config.json')


def _get_fernet():
    key = os.environ.get('SMTP_CONFIG_KEY')
    if not key:
        return None
    return Fernet(key.encode() if isinstance(key, str) else key)


def load_smtp_config():
    if not os.path.exists(SMTP_CONFIG_PATH):
        return {}
    with open(SMTP_CONFIG_PATH, 'rb') as f:
        data = f.read()
    fernet = _get_fernet()
    if fernet:
        try:
            plain = fernet.decrypt(data)
        except InvalidToken:
            return {}
        return json.loads(plain.decode('utf-8'))
    try:
        return json.loads(data.decode('utf-8'))
    except Exception:
        return {}


def save_smtp_config(cfg: dict):
    fernet = _get_fernet()
    raw = json.dumps(cfg).encode('utf-8')
    out = fernet.encrypt(raw) if fernet else raw
    with open(SMTP_CONFIG_PATH, 'wb') as f:
        f.write(out)


# ── Background match task ───────────────────────────────────────────────────────

def _run_match_task(task_id: str, search_settings: dict = None, user_id=None):
    """
    search_settings keys (all optional, fall back to parsed profile defaults):
      queries       – list of search query strings
      extra_queries – list of additional query strings to append
      min_salary    – int, minimum salary filter (default 100_000)
      remote_pref   – 'remote_only' | 'dc_only' | 'both' | 'any'
      watchlist     – list of {name, url} dicts for custom company scraping
    """
    settings = search_settings or {}
    try:
        upload_folder = get_user_upload_folder(user_id)
        resume_file = next(
            (f for f in os.listdir(upload_folder) if f.startswith('resume_')), None
        )
        if not resume_file:
            MATCH_TASKS[task_id] = {'status': 'error', 'error': 'No resume found. Please upload one first.'}
            return

        MATCH_TASKS[task_id]['status'] = 'analyzing'
        parsed = load_parsed_profile(user_id)
        if not parsed:
            resume_text = load_resume_text(user_id)
            if not resume_text:
                resume_text = extract_text(os.path.join(upload_folder, resume_file))
                with open(get_resume_text_path(user_id), 'w', encoding='utf-8') as tf:
                    tf.write(resume_text)

            cover_text = ''
            cover_file = next(
                (f for f in os.listdir(upload_folder) if f.startswith('cover_template_')), None
            )
            if cover_file:
                try:
                    cover_text = extract_text(os.path.join(upload_folder, cover_file))
                except Exception:
                    pass

            parsed = analyze_resume(resume_text, cover_text)
            save_parsed_profile(parsed, user_id)

        MATCH_TASKS[task_id]['status'] = 'searching'

        # Use user-confirmed queries if provided, else fall back to parsed profile
        queries = settings.get('queries') or parsed.get('search_queries') or []
        extra  = settings.get('extra_queries') or []
        if not queries:
            queries = [
                'data analyst', 'data engineer', 'sql developer',
                'business intelligence', 'analytics engineer',
                'tableau developer', 'databricks',
                'political data analyst', 'government technology analyst',
                'civic tech', 'campaign analytics',
            ]
        queries = list(dict.fromkeys(queries + extra))  # merge, preserve order, dedup

        # If user described their ideal position, use Gemini to generate additional targeted queries
        ideal_position = settings.get('ideal_position', '')
        if ideal_position and os.environ.get('GEMINI_API_KEY'):
            try:
                from cover_letter_generator import generate_search_queries
                ideal_queries = generate_search_queries(ideal_position, parsed)
                queries = list(dict.fromkeys(queries + ideal_queries))
            except Exception:
                pass

        watchlist = settings.get('watchlist') or []
        all_jobs = fetch_all_jobs(queries, watchlist=watchlist)

        if not all_jobs:
            MATCH_TASKS[task_id] = {'status': 'error',
                                     'error': 'No jobs returned from any source. Check your internet connection.'}
            return

        MATCH_TASKS[task_id]['status'] = 'filtering'
        # Fill missing salaries from comparable jobs before filtering
        estimate_salaries_from_comparables(list(all_jobs.values()))
        min_salary = int(settings.get('min_salary') or 100_000)
        filtered = apply_filters(list(all_jobs.values()), min_salary=min_salary)
        if not filtered:
            MATCH_TASKS[task_id] = {
                'status': 'error',
                'error': 'No jobs passed the filters ($100K+ salary, remote or DC metro/Miami/Austin).'
            }
            return

        MATCH_TASKS[task_id]['status'] = 'scoring'
        # If the user toggled specific skills on the settings page, use those for scoring
        selected_skills = settings.get('selected_skills') or []
        if selected_skills:
            parsed = dict(parsed)
            parsed['skills'] = selected_skills
        scored = score_and_categorize_jobs(filtered, parsed)
        ranked = rank_jobs(scored)

        cache_key = str(uuid.uuid4())
        JOB_CACHE[cache_key] = ranked

        MATCH_TASKS[task_id] = {
            'status': 'done',
            'cache_key': cache_key,
            'profile': parsed,
            'count': len(ranked),
        }
    except Exception as e:
        MATCH_TASKS[task_id] = {'status': 'error', 'error': str(e)}


# ── Batch apply helpers ─────────────────────────────────────────────────────────

def _run_batch(batch_id: str, profile: dict, resume_path: str):
    """Process selected jobs sequentially; waits for completion of each before next."""
    batch_session = BATCH_SESSIONS.get(batch_id)
    if not batch_session:
        return
    for i, job in enumerate(batch_session['jobs']):
        batch_session['current'] = i
        job_url = job.get('url', '')
        if not job_url:
            batch_session['statuses'][i] = 'skipped'
            continue
        app_id = f"{batch_id}_{i}"
        start_fill(app_id, job_url, profile, resume_path)
        # Wait until user confirms/cancels/errors this job before opening the next browser
        entry = PENDING_APPLICATIONS.get(app_id, {})
        entry.get('completion_event', threading.Event()).wait(timeout=3600)
        final_status = PENDING_APPLICATIONS.get(app_id, {}).get('status', 'error')
        batch_session['statuses'][i] = final_status
        # Save submitted jobs to the applications tracker
        if final_status == 'submitted':
            save_application_record(job, status='applied')
    batch_session['current'] = -1
    batch_session['done'] = True


def _run_batch_quick(batch_id: str, profile: dict, resume_path: str, max_concurrent: int = 4):
    """Process all jobs in parallel with headless auto-submit.
    No browser windows or user confirmation — fills and submits automatically.
    At most max_concurrent browsers run simultaneously to avoid resource exhaustion."""
    batch_session = BATCH_SESSIONS.get(batch_id)
    if not batch_session:
        return

    sem = threading.Semaphore(max_concurrent)

    def _apply_one(i, job):
        job_url = job.get('url', '')
        if not job_url:
            batch_session['statuses'][i] = 'skipped'
            return
        app_id = f"{batch_id}_{i}"
        sem.acquire()
        try:
            start_fill(app_id, job_url, profile, resume_path, headless=True)
            entry = PENDING_APPLICATIONS.get(app_id, {})
            entry.get('completion_event', threading.Event()).wait(timeout=300)
        finally:
            sem.release()
        final_status = PENDING_APPLICATIONS.get(app_id, {}).get('status', 'error')
        batch_session['statuses'][i] = final_status
        if final_status == 'submitted':
            save_application_record(job, status='applied')

    threads = [
        threading.Thread(target=_apply_one, args=(i, job), daemon=True)
        for i, job in enumerate(batch_session['jobs'])
    ]
    for t in threads:
        t.start()
    for t in threads:
        t.join(timeout=3600)

    batch_session['current'] = -1
    batch_session['done'] = True


# ── Routes ──────────────────────────────────────────────────────────────────────

@app.route('/')
def index():
    profile = load_profile()
    user_id = _current_user_id()
    upload_folder = get_user_upload_folder(user_id)
    files = os.listdir(upload_folder)
    resume = next((f for f in files if f.lower().startswith('resume_')), None)
    template = next((f for f in files if f.lower().startswith('cover_template')), None)
    app_counts = {}
    try:
        conn = sqlite3.connect(DB_PATH)
        c = conn.cursor()
        c.execute('SELECT status, COUNT(*) FROM applications WHERE user_id=? GROUP BY status',
                  (user_id or 0,))
        app_counts = dict(c.fetchall())
        conn.close()
    except Exception:
        pass
    return render_template('index.html', profile=profile, resume=resume,
                           template=template, app_counts=app_counts,
                           username=session.get('username', ''))


@app.route('/upload', methods=['POST'])
def upload():
    profile = {
        'name':            request.form.get('name', '').strip(),
        'email':           request.form.get('email', '').strip(),
        'phone':           request.form.get('phone', '').strip(),
        'intro':           request.form.get('intro', '').strip(),
        'skills':          request.form.get('skills', '').strip(),
        'linkedin':        request.form.get('linkedin', '').strip(),
        'website':         request.form.get('website', '').strip(),
        'city':            request.form.get('city', '').strip(),
        'ideal_position':  request.form.get('ideal_position', '').strip(),
    }
    save_profile(profile)

    user_id = _current_user_id()
    upload_folder = get_user_upload_folder(user_id)
    resume_text_path = get_resume_text_path(user_id)
    parsed_profile_path = get_parsed_profile_path(user_id)
    resume_uploaded = False
    for key in ('resume', 'cover_template'):
        f = request.files.get(key)
        if f and f.filename:
            filename = secure_filename(f.filename)
            prefix = 'resume' if key == 'resume' else 'cover_template'
            dest = os.path.join(upload_folder, f"{prefix}_{filename}")
            f.save(dest)
            if key == 'resume':
                resume_uploaded = True
                try:
                    text = extract_text(dest)
                    with open(resume_text_path, 'w', encoding='utf-8') as tf:
                        tf.write(text)
                    if os.path.exists(parsed_profile_path):
                        os.remove(parsed_profile_path)
                except Exception as e:
                    flash(f'Warning: could not extract resume text: {e}')

    # Auto-fill blank profile fields from resume if Gemini is available
    if resume_uploaded and os.environ.get('GEMINI_API_KEY'):
        resume_text = load_resume_text(user_id)
        if resume_text:
            try:
                parsed = analyze_resume(resume_text)
                save_parsed_profile(parsed, user_id)
                # Map parsed fields → profile fields, only filling blanks
                field_map = {
                    'name':    parsed.get('name', ''),
                    'email':   parsed.get('email', ''),
                    'phone':   parsed.get('phone', ''),
                    'city':    parsed.get('city', ''),
                    'linkedin': parsed.get('linkedin', ''),
                    'website': parsed.get('website', ''),
                    'intro':   parsed.get('intro', ''),
                    'skills':  parsed.get('skills_summary', ''),
                }
                filled_fields = []
                for field, value in field_map.items():
                    if value and not profile.get(field):
                        profile[field] = value
                        filled_fields.append(field)
                if filled_fields:
                    save_profile(profile)
                    flash(f'Profile auto-filled from resume: {", ".join(filled_fields)}. Review and adjust below.')
                else:
                    flash('Resume uploaded and parsed. Profile fields already filled — no changes made.')
            except Exception as e:
                flash(f'Resume uploaded. Auto-fill failed: {e}')
        else:
            flash('Profile and files saved.')
    else:
        flash('Profile and files saved.')

    # After a fresh resume upload with Gemini available, go straight to search settings
    if resume_uploaded and os.environ.get('GEMINI_API_KEY'):
        return redirect(url_for('search_settings'))
    return redirect(url_for('index'))


@app.route('/parse-resume', methods=['POST'])
def parse_resume_route():
    """Re-run Gemini analysis on the uploaded resume and overwrite profile fields."""
    if not os.environ.get('GEMINI_API_KEY'):
        flash('GEMINI_API_KEY not set — cannot auto-parse.')
        return redirect(url_for('index'))
    user_id = _current_user_id()
    resume_text = load_resume_text(user_id)
    if not resume_text:
        flash('No resume found. Please upload one first.')
        return redirect(url_for('index'))
    try:
        parsed = analyze_resume(resume_text)
        save_parsed_profile(parsed, user_id)
        profile = load_profile(user_id)
        field_map = {
            'name':    parsed.get('name', ''),
            'email':   parsed.get('email', ''),
            'phone':   parsed.get('phone', ''),
            'city':    parsed.get('city', ''),
            'linkedin': parsed.get('linkedin', ''),
            'website': parsed.get('website', ''),
            'intro':   parsed.get('intro', ''),
            'skills':  parsed.get('skills_summary', ''),
        }
        updated = []
        for field, value in field_map.items():
            if value:
                profile[field] = value
                updated.append(field)
        save_profile(profile, user_id)
        flash(f'Profile updated from resume: {", ".join(updated)}. Review below.')
    except Exception as e:
        flash(f'Parse failed: {e}')
    return redirect(url_for('index'))


# ── Job search (manual) ─────────────────────────────────────────────────────────

@app.route('/search', methods=['GET'])
def search():
    q = request.args.get('q', '')
    if not q:
        flash('Please enter a search query.')
        return redirect(url_for('index'))
    jobs = {}
    try:
        resp = requests.get(
            'https://www.themuse.com/api/public/jobs',
            params={'name': q, 'page': 0},
            timeout=20, headers={'User-Agent': 'Mozilla/5.0'}
        )
        resp.raise_for_status()
        for j in resp.json().get('results', []):
            url = (j.get('refs') or {}).get('landing_page', '')
            if not url:
                continue
            locs = [l.get('name', '') for l in (j.get('locations') or [])]
            jobs[url] = {
                'url': url,
                'title': j.get('name'),
                'company_name': (j.get('company') or {}).get('name'),
                'description': j.get('contents', ''),
                'location': ', '.join(locs),
                'remote': not locs or any('remote' in l.lower() or 'flexible' in l.lower() for l in locs),
                'salary': None,
                'date_posted': j.get('publication_date'),
            }
    except Exception:
        pass
    jobs_list = list(jobs.values())
    key = str(uuid.uuid4())
    JOB_CACHE[key] = jobs_list
    return render_template('jobs.html', jobs=jobs_list, key=key, query=q)


@app.route('/job/<cache_key>/<int:idx>')
def job_detail(cache_key, idx):
    jobs = JOB_CACHE.get(cache_key)
    if not jobs or idx < 0 or idx >= len(jobs):
        flash('Job not found in cache.')
        return redirect(url_for('index'))
    return render_template('job_detail.html', job=jobs[idx], cache_key=cache_key, idx=idx)


# ── AI matching ─────────────────────────────────────────────────────────────────

@app.route('/find-matches')
def find_matches():
    """Redirect to search settings page for confirmation before launching."""
    if not os.environ.get('GEMINI_API_KEY'):
        flash('GEMINI_API_KEY is not set in your .env file.')
        return redirect(url_for('index'))
    upload_folder = get_user_upload_folder(_current_user_id())
    resume_file = next(
        (f for f in os.listdir(upload_folder) if f.startswith('resume_')), None
    )
    if not resume_file:
        flash('Please upload a resume first.')
        return redirect(url_for('index'))
    return redirect(url_for('search_settings'))


@app.route('/search-settings')
def search_settings():
    """Show parsed resume summary and let user confirm/edit search settings."""
    if not os.environ.get('GEMINI_API_KEY'):
        flash('GEMINI_API_KEY is not set in your .env file.')
        return redirect(url_for('index'))
    user_id = _current_user_id()
    upload_folder = get_user_upload_folder(user_id)
    resume_file = next((f for f in os.listdir(upload_folder) if f.startswith('resume_')), None)
    if not resume_file:
        flash('Please upload a resume first.')
        return redirect(url_for('index'))

    parsed = load_parsed_profile(user_id)
    if not parsed:
        # Run analysis now so the page has data to show
        resume_text = load_resume_text(user_id)
        if not resume_text:
            try:
                resume_text = extract_text(os.path.join(upload_folder, resume_file))
                with open(get_resume_text_path(user_id), 'w', encoding='utf-8') as tf:
                    tf.write(resume_text)
            except Exception as e:
                flash(f'Could not read resume: {e}')
                return redirect(url_for('index'))
        try:
            parsed = analyze_resume(resume_text)
            save_parsed_profile(parsed, user_id)
            # Also pre-fill any blank profile fields
            profile = load_profile(user_id)
            field_map = {
                'name': parsed.get('name', ''), 'email': parsed.get('email', ''),
                'phone': parsed.get('phone', ''), 'city': parsed.get('city', ''),
                'linkedin': parsed.get('linkedin', ''), 'website': parsed.get('website', ''),
                'intro': parsed.get('intro', ''), 'skills': parsed.get('skills_summary', ''),
            }
            changed = False
            for field, value in field_map.items():
                if value and not profile.get(field):
                    profile[field] = value
                    changed = True
            if changed:
                save_profile(profile, user_id)
        except Exception as e:
            flash(f'Resume analysis failed: {e}')
            return redirect(url_for('index'))

    profile = load_profile(user_id)
    return render_template('search_settings.html', parsed=parsed, profile=profile,
                           cache_age=cache_age_hours())


@app.route('/clear-source-cache', methods=['POST'])
def clear_cache_route():
    """Clear the ATS disk cache so the next search fetches fresh data."""
    clear_source_cache()
    flash('Source cache cleared — next search will fetch fresh data.')
    return redirect(url_for('search_settings'))


@app.route('/launch-search', methods=['POST'])
def launch_search():
    """Accept confirmed search settings and start the matching task."""
    # Collect queries: pre-populated list + any extras the user added
    queries = [q.strip() for q in request.form.getlist('queries') if q.strip()]
    extra_raw = request.form.get('extra_queries', '')
    extra_queries = [q.strip() for q in extra_raw.replace(',', '\n').splitlines() if q.strip()]
    try:
        min_salary = int(request.form.get('min_salary') or 100_000)
    except (ValueError, TypeError):
        min_salary = 100_000

    # Parse company watchlist: "Company Name | https://url" or just "Company Name"
    watchlist_raw = request.form.get('watchlist', '')
    watchlist = []
    for line in watchlist_raw.splitlines():
        line = line.strip()
        if not line:
            continue
        if '|' in line:
            parts = line.split('|', 1)
            watchlist.append({'name': parts[0].strip(), 'url': parts[1].strip()})
        else:
            watchlist.append({'name': line, 'url': ''})

    # Save watchlist back to parsed profile for next session
    if watchlist:
        _uid = _current_user_id()
        parsed = load_parsed_profile(_uid)
        parsed['custom_companies'] = watchlist
        save_parsed_profile(parsed, _uid)

    ideal_position = request.form.get('ideal_position', '').strip()
    selected_skills = [s.strip() for s in request.form.getlist('selected_skills') if s.strip()]

    search_settings = {
        'queries':         queries,
        'extra_queries':   extra_queries,
        'min_salary':      min_salary,
        'watchlist':       watchlist,
        'ideal_position':  ideal_position,
        'selected_skills': selected_skills,
    }

    user_id = _current_user_id()
    task_id = str(uuid.uuid4())
    MATCH_TASKS[task_id] = {'status': 'starting'}
    threading.Thread(target=_run_match_task, args=(task_id, search_settings, user_id),
                     daemon=True).start()
    return render_template('loading.html', task_id=task_id)


@app.route('/match-status/<task_id>')
def match_status(task_id):
    task = MATCH_TASKS.get(task_id, {'status': 'not_found'})
    resp = {'status': task.get('status'), 'error': task.get('error', '')}
    if task.get('status') == 'done':
        resp['redirect'] = url_for('match_results', key=task['cache_key'])
        resp['count'] = task.get('count', 0)
    return jsonify(resp)


@app.route('/match-results/<key>')
def match_results(key):
    jobs = JOB_CACHE.get(key, [])
    if not jobs:
        flash('Results expired. Please run matching again.')
        return redirect(url_for('index'))

    # Apply client-requested filters
    cat_filter  = request.args.get('category', '')
    company_filter = request.args.get('company', '').strip().lower()
    remote_only = request.args.get('remote_only', '')
    pol_only    = request.args.get('pol_only', '')
    dc_only     = request.args.get('dc_only', '')
    min_score   = int(request.args.get('min_score', 0))
    min_pay_str = request.args.get('min_pay', '0')
    try:
        min_pay = int(min_pay_str)
    except ValueError:
        min_pay = 0

    filtered = jobs
    if cat_filter:
        filtered = [j for j in filtered if j.get('category') == cat_filter]
    if company_filter:
        filtered = [j for j in filtered if company_filter in
                    (j.get('company_name') or j.get('company') or '').lower()]
    if remote_only == '1':
        filtered = [j for j in filtered if j.get('is_remote')]
    if pol_only == '1':
        filtered = [j for j in filtered if j.get('political_tech')]
    if dc_only == '1':
        filtered = [j for j in filtered if j.get('is_dc_metro')]
    if min_score:
        filtered = [j for j in filtered if j.get('ranking_score', 0) >= min_score]
    if min_pay:
        filtered = [j for j in filtered if (j.get('salary_max') or 0) >= min_pay]

    # Attach original index so the template can link to the right job in the full list
    orig_idx_map = {id(j): i for i, j in enumerate(jobs)}
    for j in filtered:
        j['_orig_idx'] = orig_idx_map.get(id(j), 0)

    profile = load_parsed_profile()
    return render_template(
        'matched_jobs.html',
        jobs=filtered, key=key,
        profile=profile, categories=CATEGORIES,
        filters={
            'category':   cat_filter,
            'company':    company_filter,
            'remote_only': remote_only,
            'pol_only':   pol_only,
            'dc_only':    dc_only,
            'min_score':  min_score,
            'min_pay':    min_pay,
        }
    )


# ── Cover letter ────────────────────────────────────────────────────────────────

def _get_base_letter(profile: dict, user_id=None) -> str:
    upload_folder = get_user_upload_folder(user_id or _current_user_id())
    template_file = next(
        (f for f in os.listdir(upload_folder) if f.startswith('cover_template_')), None
    )
    if template_file:
        path = os.path.join(upload_folder, template_file)
        ext = os.path.splitext(template_file)[1].lower()
        try:
            if ext == '.docx':
                from docx import Document
                doc = Document(path)
                return '\n'.join(p.text for p in doc.paragraphs)
            elif ext == '.pdf':
                from resume_parser import extract_text
                return extract_text(path)
            else:
                with open(path, 'r', encoding='utf-8', errors='ignore') as f:
                    return f.read()
        except Exception:
            pass
    return DEFAULT_TEMPLATE.format(
        company='', title='', name=profile.get('name', ''),
        intro=profile.get('intro', ''), skills=profile.get('skills', ''),
    )


@app.route('/preview/<cache_key>/<int:idx>', methods=['GET', 'POST'])
def preview(cache_key, idx):
    jobs = JOB_CACHE.get(cache_key)
    if not jobs or idx < 0 or idx >= len(jobs):
        flash('Job not found.')
        return redirect(url_for('index'))
    job = jobs[idx]
    user_id = _current_user_id()
    profile = load_profile(user_id)

    base = _get_base_letter(profile, user_id)
    try:
        filled = base.format(
            company=job.get('company_name', ''),
            title=job.get('title', ''),
            name=profile.get('name', ''),
            intro=profile.get('intro', ''),
            skills=profile.get('skills', ''),
        )
    except KeyError:
        filled = base  # template may not have all placeholders

    if request.method == 'POST':
        action = request.form.get('action', 'package')
        edited = request.form.get('edited_letter', '')
        upload_folder = get_user_upload_folder(user_id)
        resume_file = next(
            (f for f in os.listdir(upload_folder) if f.startswith('resume_')), None
        )
        if action == 'package':
            mem = io.BytesIO()
            with zipfile.ZipFile(mem, 'w') as z:
                if resume_file:
                    z.write(os.path.join(upload_folder, resume_file),
                            arcname='resume' + os.path.splitext(resume_file)[1])
                z.writestr('cover_letter.txt', edited)
            mem.seek(0)
            return send_file(mem, mimetype='application/zip',
                             as_attachment=True, download_name='application_package.zip')
        elif action == 'save':
            save_application_record(job, status='saved', cover_letter=edited)
            flash('Job saved to your applications tracker.')
            return redirect(url_for('saved_jobs'))

    return render_template('preview.html', job=job, filled=filled,
                           cache_key=cache_key, idx=idx)


@app.route('/resume-feedback/<cache_key>/<int:idx>')
def resume_feedback(cache_key, idx):
    jobs = JOB_CACHE.get(cache_key)
    if not jobs or idx < 0 or idx >= len(jobs):
        return jsonify({'error': 'Job not found'}), 404
    job = jobs[idx]
    resume_text = load_resume_text(_current_user_id())
    if not resume_text:
        return jsonify({'error': 'No resume text found. Please upload a resume.'}), 400
    try:
        feedback = generate_resume_feedback(job, resume_text)
        return jsonify(feedback)
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/tailor-letter/<cache_key>/<int:idx>')
def tailor_letter(cache_key, idx):
    jobs = JOB_CACHE.get(cache_key)
    if not jobs or idx < 0 or idx >= len(jobs):
        return jsonify({'error': 'Job not found'}), 404
    job = jobs[idx]
    user_id = _current_user_id()
    profile = load_profile(user_id)
    resume_text = load_resume_text(user_id)
    if not resume_text:
        return jsonify({'error': 'No resume text found. Please upload a resume first.'}), 400
    base = _get_base_letter(profile, user_id)
    try:
        tailored = tailor_cover_letter(job, resume_text, base, profile)
        return jsonify({'letter': tailored})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/discover-companies', methods=['POST'])
def discover_companies():
    """Use Gemini to suggest companies in a city that hire for the user's profile."""
    if not os.environ.get('GEMINI_API_KEY'):
        return jsonify({'error': 'GEMINI_API_KEY not set.'}), 400
    data = request.get_json() or {}
    city = data.get('city', '').strip()
    if not city:
        return jsonify({'error': 'City is required.'}), 400
    user_id = _current_user_id()
    parsed = load_parsed_profile(user_id)
    skills = parsed.get('skills', [])
    target_roles = parsed.get('target_roles', [])
    try:
        from job_sources import discover_companies_in_city
        companies = discover_companies_in_city(city, skills, target_roles)
        # Cache city discoveries to profile for future reference
        city_key = city.lower().strip()
        parsed.setdefault('city_discoveries', {})[city_key] = companies
        save_parsed_profile(parsed, user_id)
        return jsonify({'companies': companies, 'city': city})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


# ── Auto-fill ───────────────────────────────────────────────────────────────────

@app.route('/start-apply/<cache_key>/<int:idx>', methods=['POST'])
def start_apply(cache_key, idx):
    jobs = JOB_CACHE.get(cache_key)
    if not jobs or idx < 0 or idx >= len(jobs):
        app.logger.warning('start_apply: cache_key=%s idx=%d not found', cache_key, idx)
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'error': 'Job results not found — please run a new search.'}), 404
        flash('Job results not found — please run a new search.')
        return redirect(url_for('index'))
    job = jobs[idx]
    job_url = job.get('url')
    if not job_url:
        app.logger.warning('start_apply: no url for cache_key=%s idx=%d', cache_key, idx)
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'error': 'No application URL for this job.'}), 400
        flash('No application URL for this job.')
        return redirect(url_for('preview', cache_key=cache_key, idx=idx))

    user_id = _current_user_id()
    upload_folder = get_user_upload_folder(user_id)
    resume_file = next(
        (f for f in os.listdir(upload_folder) if f.startswith('resume_')), None
    )
    resume_path = os.path.join(upload_folder, resume_file) if resume_file else None
    profile = load_profile(user_id)
    resume_text = load_resume_text(user_id)

    # Include cover letter from the preview form if the user edited it
    cover_letter = request.form.get('edited_letter', '')
    if not cover_letter:
        try:
            base = _get_base_letter(profile, user_id)
            cover_letter = base.format(
                company=job.get('company_name', ''),
                title=job.get('title', ''),
                name=profile.get('name', ''),
                intro=profile.get('intro', ''),
                skills=profile.get('skills', ''),
            )
        except Exception:
            cover_letter = ''

    full_profile = {
        **profile,
        'cover_letter':       cover_letter,
        'work_auth':          request.form.get('work_auth', profile.get('work_auth', 'yes')),
        'sponsorship':        request.form.get('sponsorship', profile.get('sponsorship', 'no')),
        'salary_expectation': request.form.get('salary_expectation', profile.get('salary_expectation', '')),
        # Context for AI custom question answering
        'resume_text':        resume_text,
        '_job_title':         job.get('title', ''),
        '_company':           job.get('company_name') or job.get('company', ''),
        '_job_description':   (job.get('description') or '')[:3000],
    }

    app_id = str(uuid.uuid4())
    app.logger.info('start_apply: app_id=%s url=%s resume=%s', app_id, job_url, resume_path)
    start_fill(app_id, job_url, full_profile, resume_path)
    return render_template('apply_confirm.html', app_id=app_id, job=job,
                           cache_key=cache_key, idx=idx)


@app.route('/quick-apply/<cache_key>/<int:idx>', methods=['POST'])
def quick_apply(cache_key, idx):
    """Headless Quick Apply: invisible browser fills and submits; auto-tracked."""
    jobs = JOB_CACHE.get(cache_key)
    if not jobs or idx < 0 or idx >= len(jobs):
        flash('Job not found.')
        return redirect(url_for('index'))
    job = jobs[idx]
    job_url = job.get('url')
    if not job_url:
        flash('No application URL for this job.')
        return redirect(url_for('match_results', key=cache_key))

    user_id = _current_user_id()
    profile = load_profile(user_id)
    resume_text = load_resume_text(user_id)
    upload_folder = get_user_upload_folder(user_id)
    resume_file = next(
        (f for f in os.listdir(upload_folder) if f.startswith('resume_')), None
    )
    resume_path = os.path.join(upload_folder, resume_file) if resume_file else None

    extra_info = {
        'work_auth':          request.form.get('work_auth', profile.get('work_auth', 'yes')),
        'sponsorship':        request.form.get('sponsorship', profile.get('sponsorship', 'no')),
        'salary_expectation': request.form.get('salary_expectation', profile.get('salary_expectation', '')),
        'start_date':         request.form.get('start_date', ''),
        'referral':           request.form.get('referral', ''),
    }

    # Generate AI-tailored cover letter
    cover_letter = ''
    if os.environ.get('GEMINI_API_KEY') and resume_text:
        try:
            base = _get_base_letter(profile, user_id)
            cover_letter = tailor_cover_letter(job, resume_text, base, profile)
        except Exception:
            pass
    extra_info['cover_letter'] = cover_letter

    full_profile = {
        **profile,
        **extra_info,
        # Context for AI custom question answering
        'resume_text':      resume_text,
        '_job_title':       job.get('title', ''),
        '_company':         job.get('company_name') or job.get('company', ''),
        '_job_description': (job.get('description') or '')[:3000],
    }
    app_id = str(uuid.uuid4())

    # Save to tracker immediately as 'pending', update to 'applied' on success
    job_record = dict(job)
    save_application_record(job_record, status='pending', cover_letter=cover_letter,
                            user_id=user_id)

    def _on_complete(entry):
        status = entry.get('status', 'error')
        if status == 'submitted':
            save_application_record(job_record, status='applied',
                                    cover_letter=cover_letter, user_id=user_id)

    start_fill(app_id, job_url, full_profile, resume_path,
               headless=True, on_complete=_on_complete)

    return render_template('quick_apply_status.html', app_id=app_id, job=job,
                           cache_key=cache_key, idx=idx)


@app.route('/apply-status/<app_id>')
def apply_status_route(app_id):
    entry = PENDING_APPLICATIONS.get(app_id, {})
    return jsonify({
        'status':       entry.get('status', 'not_found'),
        'message':      entry.get('message', ''),
        'filled_fields': entry.get('filled_fields', []),
        'error':        entry.get('error', ''),
        'ats':          entry.get('ats', ''),
        'job_url':      entry.get('job_url', ''),
    })


@app.route('/debug')
def debug_state():
    """Show live state of job cache, pending applications, and app config."""
    import _io
    cache_keys = list(JOB_CACHE.keys())
    disk_keys = []
    try:
        disk_keys = [f[:-5] for f in os.listdir(_JOB_CACHE_DIR) if f.endswith('.json')]
    except Exception:
        pass
    pending = {k: {kk: vv for kk, vv in v.items() if kk not in ('event', 'completion_event')}
               for k, v in PENDING_APPLICATIONS.items()}
    upload_files = os.listdir(app.config['UPLOAD_FOLDER'])
    return jsonify({
        'job_cache_keys_memory': cache_keys,
        'job_cache_keys_disk': disk_keys,
        'pending_applications': pending,
        'upload_files': upload_files,
        'profile_name': load_profile().get('name'),
        'resume_text_exists': os.path.exists(RESUME_TEXT_PATH),
    })


@app.route('/test-browser')
def test_browser():
    """Quick smoke-test: opens a browser to example.com and returns status."""
    import uuid as _uuid
    profile = load_profile()
    test_id = str(_uuid.uuid4())
    start_fill(test_id, 'https://example.com', profile, None)
    import time as _time
    for _ in range(12):
        _time.sleep(1)
        entry = PENDING_APPLICATIONS.get(test_id, {})
        status = entry.get('status', 'unknown')
        if status not in ('starting',):
            return jsonify({'ok': True, 'status': status, 'error': entry.get('error', '')})
    return jsonify({'ok': False, 'status': 'timeout', 'error': 'Browser did not start within 12s'})


@app.route('/confirm-apply/<app_id>', methods=['POST'])
def confirm_apply(app_id):
    entry = PENDING_APPLICATIONS.get(app_id)
    if not entry:
        return jsonify({'ok': False, 'error': 'Session expired.'}), 404
    entry['confirmed'] = True
    entry['event'].set()
    # Save to tracker (the autofiller will set status='submitted' async)
    job_url = entry.get('job_url', '')
    if job_url:
        # Find job in cache to get full metadata, or construct minimal record
        job = {'url': job_url, 'title': '', 'company_name': ''}
        for v in JOB_CACHE.values():
            if isinstance(v, list):
                for j in v:
                    if j.get('url') == job_url:
                        job = j
                        break
        save_application_record(job, status='applied', cover_letter='')
    return jsonify({'ok': True})


@app.route('/cancel-apply/<app_id>', methods=['POST'])
def cancel_apply(app_id):
    entry = PENDING_APPLICATIONS.get(app_id)
    if entry:
        entry['confirmed'] = False
        entry['event'].set()
    return jsonify({'ok': True})


@app.route('/login-done/<app_id>', methods=['POST'])
def login_done(app_id):
    """Signal that the user has logged in and the filler thread can resume."""
    entry = PENDING_APPLICATIONS.get(app_id)
    if not entry:
        return jsonify({'ok': False, 'error': 'Session expired.'}), 404
    entry['login_cancelled'] = False
    entry['login_event'].set()
    return jsonify({'ok': True})


@app.route('/login-cancel/<app_id>', methods=['POST'])
def login_cancel(app_id):
    """Cancel an application that is waiting for login."""
    entry = PENDING_APPLICATIONS.get(app_id)
    if not entry:
        return jsonify({'ok': False, 'error': 'Session expired.'}), 404
    entry['login_cancelled'] = True
    entry['login_event'].set()
    return jsonify({'ok': True})


@app.route('/batch-login-done/<batch_id>/<int:job_idx>', methods=['POST'])
def batch_login_done(batch_id, job_idx):
    """Signal login complete for a batched job."""
    app_id = f"{batch_id}_{job_idx}"
    entry  = PENDING_APPLICATIONS.get(app_id)
    if not entry:
        return jsonify({'error': 'not found'}), 404
    entry['login_cancelled'] = False
    entry['login_event'].set()
    return jsonify({'ok': True})


# ── Batch apply ─────────────────────────────────────────────────────────────────

@app.route('/prepare-batch', methods=['POST'])
def prepare_batch():
    cache_key = request.form.get('cache_key', '')
    jobs_all  = JOB_CACHE.get(cache_key, [])
    indices   = request.form.getlist('selected_jobs')
    if not indices:
        flash('No jobs selected.')
        return redirect(url_for('match_results', key=cache_key))
    selected = []
    for idx_str in indices:
        try:
            idx = int(idx_str)
            if 0 <= idx < len(jobs_all):
                selected.append(dict(jobs_all[idx]))  # copy so we can annotate
        except ValueError:
            pass
    if not selected:
        flash('No valid jobs selected.')
        return redirect(url_for('match_results', key=cache_key))
    batch_id = str(uuid.uuid4())
    BATCH_SESSIONS[batch_id] = {
        'jobs':      selected,
        'cache_key': cache_key,
        'extra_info': {},
        'statuses':  {i: 'pending' for i in range(len(selected))},
        'current':   -1,
        'done':      False,
    }
    return redirect(url_for('apply_prep', batch_id=batch_id))


@app.route('/apply-prep/<batch_id>')
def apply_prep(batch_id):
    batch_session = BATCH_SESSIONS.get(batch_id)
    if not batch_session:
        flash('Session expired.')
        return redirect(url_for('index'))
    user_id = _current_user_id()
    profile = load_profile(user_id)
    parsed  = load_parsed_profile(user_id)
    upload_folder = get_user_upload_folder(user_id)
    cover_text = ''
    template_file = next(
        (f for f in os.listdir(upload_folder) if f.startswith('cover_template_')), None
    )
    if template_file:
        try:
            with open(os.path.join(upload_folder, template_file),
                      'r', encoding='utf-8', errors='ignore') as f:
                cover_text = f.read()
        except Exception:
            pass
    return render_template('apply_prep.html',
                           batch_id=batch_id,
                           jobs=batch_session['jobs'],
                           profile=profile,
                           parsed=parsed,
                           cover_text=cover_text)


@app.route('/launch-batch/<batch_id>', methods=['POST'])
def launch_batch(batch_id):
    batch_session = BATCH_SESSIONS.get(batch_id)
    if not batch_session:
        flash('Session expired.')
        return redirect(url_for('index'))
    # Save user-supplied extra info
    extra_info = {
        'work_auth':          request.form.get('work_auth', ''),
        'sponsorship':        request.form.get('sponsorship', ''),
        'salary_expectation': request.form.get('salary_expectation', ''),
        'start_date':         request.form.get('start_date', ''),
        'referral':           request.form.get('referral', ''),
        'cover_letter':       request.form.get('cover_letter', ''),
    }
    batch_session['extra_info'] = extra_info
    # Honour per-job de-selection on the prep page
    keep = set(request.form.getlist('apply_to'))
    batch_session['jobs']     = [j for i, j in enumerate(batch_session['jobs']) if str(i) in keep]
    batch_session['statuses'] = {i: 'pending' for i in range(len(batch_session['jobs']))}
    if not batch_session['jobs']:
        flash('No jobs selected to apply to.')
        return redirect(url_for('apply_prep', batch_id=batch_id))
    user_id = _current_user_id()
    profile     = load_profile(user_id)
    upload_folder = get_user_upload_folder(user_id)
    resume_file = next(
        (f for f in os.listdir(upload_folder) if f.startswith('resume_')), None
    )
    resume_path = os.path.join(upload_folder, resume_file) if resume_file else None
    full_profile = {**profile, **extra_info}
    mode = request.form.get('mode', 'review')
    batch_session['mode'] = mode
    if mode == 'quick':
        threading.Thread(target=_run_batch_quick, args=(batch_id, full_profile, resume_path),
                         daemon=True).start()
    else:
        threading.Thread(target=_run_batch, args=(batch_id, full_profile, resume_path),
                         daemon=True).start()
    return redirect(url_for('batch_status_page', batch_id=batch_id))


@app.route('/batch-status/<batch_id>')
def batch_status_page(batch_id):
    batch_session = BATCH_SESSIONS.get(batch_id)
    if not batch_session:
        flash('Session expired.')
        return redirect(url_for('index'))
    return render_template('batch_status.html', batch_id=batch_id,
                           jobs=batch_session['jobs'],
                           mode=batch_session.get('mode', 'review'))


@app.route('/batch-status-json/<batch_id>')
def batch_status_json(batch_id):
    batch_session = BATCH_SESSIONS.get(batch_id)
    if not batch_session:
        return jsonify({'error': 'not found'}), 404
    jobs_status = []
    for i, job in enumerate(batch_session['jobs']):
        app_id = f"{batch_id}_{i}"
        entry  = PENDING_APPLICATIONS.get(app_id, {})
        jobs_status.append({
            'idx':           i,
            'title':         job.get('title', ''),
            'company':       job.get('company_name') or job.get('company', ''),
            'url':           job.get('url', ''),
            'status':        entry.get('status') or batch_session['statuses'].get(i, 'pending'),
            'filled_fields': entry.get('filled_fields', []),
            'message':       entry.get('message', ''),
            'error':         entry.get('error', ''),
            'app_id':        app_id,
        })
    return jsonify({
        'jobs':    jobs_status,
        'current': batch_session.get('current', -1),
        'done':    batch_session.get('done', False),
        'mode':    batch_session.get('mode', 'review'),
    })


@app.route('/batch-confirm/<batch_id>/<int:job_idx>', methods=['POST'])
def batch_confirm(batch_id, job_idx):
    app_id = f"{batch_id}_{job_idx}"
    entry  = PENDING_APPLICATIONS.get(app_id)
    if not entry:
        return jsonify({'error': 'not found'}), 404
    action = request.form.get('action', 'confirm')
    entry['confirmed'] = (action == 'confirm')
    entry['event'].set()
    return jsonify({'ok': True})


# ── Application tracker ─────────────────────────────────────────────────────────

@app.route('/saved-jobs')
def saved_jobs():
    status_filter = request.args.get('status', '')
    apps = get_applications(status_filter or None)
    # Compute counts per status for the stats strip
    all_apps = get_applications()
    counts = {}
    for a in all_apps:
        s = a.get('status', 'saved')
        counts[s] = counts.get(s, 0) + 1
    return render_template('saved_jobs.html', applications=apps,
                           status_filter=status_filter, counts=counts)


@app.route('/update-application/<int:app_id>', methods=['POST'])
def update_application(app_id):
    status = request.form.get('status', 'saved')
    notes = request.form.get('notes', '')
    update_application_status(app_id, status, notes)
    flash('Application updated.')
    return redirect(url_for('saved_jobs'))


@app.route('/update-notes/<int:app_id>', methods=['POST'])
def update_notes(app_id):
    notes = request.form.get('notes', '')
    conn = sqlite3.connect(DB_PATH)
    conn.execute('UPDATE applications SET notes=? WHERE id=?', (notes, app_id))
    conn.commit()
    conn.close()
    status_filter = request.args.get('status', '')
    return redirect(url_for('saved_jobs', status=status_filter) if status_filter else url_for('saved_jobs'))


@app.route('/delete-application/<int:app_id>', methods=['POST'])
def delete_application(app_id):
    conn = sqlite3.connect(DB_PATH)
    conn.execute('DELETE FROM applications WHERE id=?', (app_id,))
    conn.commit()
    conn.close()
    flash('Application removed.')
    return redirect(url_for('saved_jobs'))


# ── Email send ──────────────────────────────────────────────────────────────────

@app.route('/send-application', methods=['POST'])
def send_application():
    recipient = request.form.get('recipient')
    subject = request.form.get('subject') or 'Application'
    body = request.form.get('edited_letter') or ''
    resume_file = next(
        (f for f in os.listdir(app.config['UPLOAD_FOLDER']) if f.startswith('resume_')), None
    )
    if not recipient:
        flash('Recipient required')
        return redirect(request.referrer or url_for('index'))

    smtp_cfg = load_smtp_config()
    smtp_host = request.form.get('smtp_host') or smtp_cfg.get('smtp_host') or os.environ.get('SMTP_HOST')
    smtp_port = int(request.form.get('smtp_port') or smtp_cfg.get('smtp_port') or os.environ.get('SMTP_PORT') or 587)
    smtp_user = request.form.get('smtp_user') or smtp_cfg.get('smtp_user') or os.environ.get('SMTP_USER')
    smtp_pass = request.form.get('smtp_pass') or smtp_cfg.get('smtp_pass') or os.environ.get('SMTP_PASS')

    msg = EmailMessage()
    msg['From'] = smtp_user or 'noreply@example.com'
    msg['To'] = recipient
    msg['Subject'] = subject
    msg.set_content(body)

    if resume_file:
        path = os.path.join(app.config['UPLOAD_FOLDER'], resume_file)
        try:
            with open(path, 'rb') as f:
                data = f.read()
            msg.add_attachment(data, maintype='application', subtype='octet-stream',
                               filename=resume_file)
        except Exception as e:
            flash(f'Failed to attach resume: {e}')

    try:
        with smtplib.SMTP(smtp_host, smtp_port, timeout=30) as server:
            server.starttls()
            if smtp_user and smtp_pass:
                server.login(smtp_user, smtp_pass)
            server.send_message(msg)
        flash('Email sent')
    except Exception as e:
        flash(f'Email failed: {e}')

    return redirect(url_for('index'))


# ── Scraped jobs DB browser ──────────────────────────────────────────────────────

@app.route('/jobs_db')
def jobs_db():
    filters = {
        'company': request.args.get('company'),
        'title': request.args.get('title'),
        'location': request.args.get('location'),
        'remote': None,
        'since_days': request.args.get('since_days') or 14,
    }
    r = request.args.get('remote')
    if r is not None:
        filters['remote'] = r.lower() in ('1', 'true', 'yes')
    jobs = query_jobs(filters)
    return render_template('jobs_db.html', jobs=jobs, filters=filters)


# ── Site credentials (job board accounts for Quick Apply auto-login) ─────────────

SITE_CREDENTIALS_PATH = os.path.join(os.path.dirname(__file__), 'site_credentials.json')

def _load_site_creds() -> dict:
    try:
        with open(SITE_CREDENTIALS_PATH, 'r', encoding='utf-8') as f:
            return json.load(f)
    except Exception:
        return {}

def _save_site_creds(creds: dict):
    with open(SITE_CREDENTIALS_PATH, 'w', encoding='utf-8') as f:
        json.dump(creds, f, indent=2)


@app.route('/site-credentials')
def site_credentials_page():
    creds = _load_site_creds()
    return render_template('site_credentials.html', creds=creds)


@app.route('/site-credentials', methods=['POST'])
def save_site_credentials():
    action  = request.form.get('action', 'save')
    domain  = request.form.get('domain', '').strip().lower()
    email   = request.form.get('email', '').strip()
    password = request.form.get('password', '').strip()

    creds = _load_site_creds()

    if action == 'delete' and domain:
        creds.pop(domain, None)
        flash(f'Removed credentials for {domain}.')
    elif action == 'save' and domain and email and password:
        creds[domain] = {'email': email, 'password': password}
        flash(f'Credentials saved for {domain}.')
    else:
        flash('Please fill in all fields.')
        return redirect(url_for('site_credentials_page'))

    _save_site_creds(creds)
    return redirect(url_for('site_credentials_page'))


# ── Entry point ─────────────────────────────────────────────────────────────────

if __name__ == '__main__':
    init_db()
    threading.Timer(1.0, lambda: webbrowser.open('http://127.0.0.1:5000')).start()
    app.run(debug=False, threaded=True)
